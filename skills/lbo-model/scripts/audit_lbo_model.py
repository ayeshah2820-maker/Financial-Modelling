from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


SHEET_PATTERNS = {
    "assumptions": r"assump|input|假设",
    "sources_uses": r"source|s&u|transaction|交易",
    "income_statement": r"income|\bis\b|利润",
    "balance_sheet": r"balance|\bbs\b|资产负债",
    "cash_flow": r"cash flow|\bcf\b|现金流",
    "debt_schedule": r"debt|债务",
    "returns": r"return|回报",
    "checks": r"check|检查",
}
ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#N/A")


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit the structure and formulas of an LBO workbook.")
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()

    wb = load_workbook(args.workbook, data_only=False, read_only=False, keep_links=True)
    cached_wb = load_workbook(args.workbook, data_only=True, read_only=False, keep_links=True)
    names = wb.sheetnames
    findings = {"workbook": str(args.workbook), "errors": [], "warnings": [], "metrics": {}}

    for role, pattern in SHEET_PATTERNS.items():
        if not any(re.search(pattern, name, re.I) for name in names):
            findings["errors"].append(f"Missing sheet role: {role}")

    formula_count = 0
    external_link_formulas = []
    formula_errors = []
    hardcoded_blue = 0
    cross_sheet_green = 0
    formula_black = 0
    text_precedent_formulas = []
    cached_errors = []

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text = str(value)
                color = getattr(cell.font.color, "rgb", None) if cell.font.color else None
                color = str(color).upper() if color else ""
                if text.startswith("="):
                    formula_count += 1
                    if any(token in text.upper() for token in ERROR_TOKENS):
                        formula_errors.append(f"{ws.title}!{cell.coordinate}: {text}")
                    if "[" in text and "]" in text:
                        external_link_formulas.append(f"{ws.title}!{cell.coordinate}: {text}")
                    if "!" in text and color.endswith("008000"):
                        cross_sheet_green += 1
                    elif color.endswith("000000") or not color:
                        formula_black += 1
                    # Catch a common helper-table bug: formulas referencing the
                    # adjacent hardcoded label cell instead of the value cell.
                    for ref in re.findall(r"(?<![!A-Z0-9_])(\$?[A-Z]{1,3}\$?\d+)", text.upper()):
                        precedent = ws[ref.replace("$", "")].value
                        if isinstance(precedent, str) and not precedent.startswith("="):
                            text_precedent_formulas.append(
                                f"{ws.title}!{cell.coordinate} references text cell {ref}: {precedent}"
                            )
                elif isinstance(value, (int, float)) and color.endswith("0000FF"):
                    hardcoded_blue += 1

    # Formula text can be syntactically valid while native Excel calculates an
    # error (for example, a formula that points to an adjacent text-label cell).
    # A data-only pass catches errors saved in the workbook's cached results.
    for ws in cached_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if cell.data_type == "e" or (isinstance(value, str) and value.upper() in ERROR_TOKENS):
                    cached_errors.append(f"{ws.title}!{cell.coordinate}: {value}")

    findings["metrics"].update({
        "sheet_count": len(names),
        "formula_count": formula_count,
        "blue_numeric_inputs": hardcoded_blue,
        "green_cross_sheet_formulas": cross_sheet_green,
        "black_formula_cells": formula_black,
        "cached_formula_errors": len(cached_errors),
    })
    if formula_count < 100:
        findings["warnings"].append("Workbook has fewer than 100 formulas; confirm it is not overly hardcoded.")
    if formula_errors:
        findings["errors"].extend(formula_errors[:50])
    if cached_errors:
        findings["errors"].append(f"Cached formula errors found: {len(cached_errors)}")
        findings["errors"].extend(cached_errors[:50])
    if external_link_formulas:
        findings["warnings"].append(f"External-link formulas found: {len(external_link_formulas)}")
        findings["warnings"].extend(external_link_formulas[:20])
    if text_precedent_formulas:
        findings["warnings"].append(
            f"Formulas referencing hardcoded text cells found: {len(text_precedent_formulas)}"
        )
        findings["warnings"].extend(text_precedent_formulas[:20])

    wb.close()
    cached_wb.close()
    findings["status"] = "FAIL" if findings["errors"] else "PASS"
    print(json.dumps(findings, ensure_ascii=False, indent=2))
    return 1 if findings["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
